"""Export one training-target row per rat from the Stanford study workbook."""

import csv
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK_PATH = Path(__file__).parent / "Stanford_Murine_Fat_Quantification" / "SoSDataAll.xlsx"
OUTPUT_PATH = Path(__file__).parent / "training_targets.csv"


# Column numbers are zero-based and correspond to the Data sheet's row-2 headers.
TARGET_COLUMNS = {
    "rat_id": 0,
    "gender": 1,
    "diet_week": 2,
    "weight_g": 3,
    "ground_truth_sos_mps": 4,
    "ground_truth_sos_std_mps": 5,
    "local_sos_fsa_mps": 14,
    "local_sos_fsa_std_mps": 15,
    "local_sos_hadamard_mps": 16,
    "local_sos_hadamard_std_mps": 17,
    "local_sos_multifocal_mps": 18,
    "local_sos_multifocal_std_mps": 19,
    "average_sos_fsa_mps": 20,
    "average_sos_fsa_std_mps": 21,
    "average_sos_hadamard_mps": 22,
    "average_sos_hadamard_std_mps": 23,
    "average_sos_multifocal_mps": 24,
    "average_sos_multifocal_std_mps": 25,
    "liver_fat_percent": 26,
    "liver_fat_std_percent": 27,
    "non_fat_percent": 28,
    "steatosis_grade": 29,
    "lipidosis_grade": 30,
    "water_temperature_c": 31,
    "sos_at_diet_average_mps": 32,
    "old_fat_average_percent": 33,
    "old_fat_std_percent": 34,
}


def clean_value(value):
    """Represent Excel's missing and textual placeholder values as empty cells."""
    if value is None or (isinstance(value, str) and value.strip().lower() in {"", "n/a", "na"}):
        return ""
    return value


def export_training_targets(workbook_path=WORKBOOK_PATH, output_path=OUTPUT_PATH):
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    sheet = workbook["Data"]
    output_fields = list(TARGET_COLUMNS)

    with output_path.open("w", newline="", encoding="utf-8") as csv_file:
        writer = csv.DictWriter(csv_file, fieldnames=output_fields)
        writer.writeheader()

        # Row 1 is a category header and row 2 contains the actual column labels.
        for row in sheet.iter_rows(min_row=3, values_only=True):
            rat_id = clean_value(row[TARGET_COLUMNS["rat_id"]])
            if rat_id == "":
                continue
            writer.writerow(
                {
                    field: clean_value(row[column_index])
                    for field, column_index in TARGET_COLUMNS.items()
                }
            )


if __name__ == "__main__":
    export_training_targets()
    print(f"Wrote training targets to {OUTPUT_PATH}")